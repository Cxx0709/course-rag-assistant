import io
import json

import pymupdf
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from course_rag_assistant.web import create_app

NOTES = "# 训练集与测试集\n\n训练集用于拟合模型参数，测试集用于评估模型的泛化能力。"


@pytest.fixture
def client(tmp_path):
    app = create_app(tmp_path / "library", generator_factory=lambda: None)
    with TestClient(app, base_url="http://127.0.0.1") as client:
        client.headers["X-Workspace-Token"] = client.get("/api/state").json()["token"]
        yield client


def upload(client, name="notes.md", text=NOTES):
    response = client.post("/api/documents", params={"filename": name}, content=text.encode("utf-8"))
    assert response.status_code == 201, response.text
    return response.json()


def ask(client, ids, question="训练集和测试集的用途有什么不同？", **kwargs):
    return client.post("/api/ask", json={"question": question, "document_ids": ids, **kwargs})


def test_upload_ask_and_cited_original(client):
    document = upload(client)
    response = ask(client, [document["id"]])
    assert response.status_code == 200
    result = response.json()
    assert not result["abstained"] and result["mode"] == "source"
    assert "拟合模型参数" in result["claims"][0]["text"]
    assert result["claims"][0]["source_ids"] == ["1"]
    source = result["sources"][0]
    assert source["name"] == "notes.md" and source["section"] == "训练集与测试集"
    assert source["text"] in NOTES
    original = client.get(source["url"])
    assert original.content == NOTES.encode("utf-8")
    assert "attachment" in original.headers["content-disposition"]
    assert original.headers["x-content-type-options"] == "nosniff"


def test_same_named_files_and_question_scope_do_not_mix_sources(client):
    a = upload(client, text="# 模型学习率\n\n模型学习率设置为0.03。")
    b = upload(client, text="# 模型学习率\n\n模型学习率设置为0.09。")
    assert a["id"] != b["id"]
    result = ask(client, [a["id"]], "模型学习率设置为多少？").json()
    assert "0.03" in json.dumps(result, ensure_ascii=False)
    assert "0.09" not in json.dumps(result, ensure_ascii=False)
    assert all(s["document_id"] == a["id"] for s in result["sources"])
    stored = client.app.state.library.documents
    assert stored[a["id"]].source != stored[b["id"]].source


def test_upload_survives_restart_and_uses_new_request_token(client):
    document = upload(client)
    app = create_app(client.app.state.library.root, generator_factory=lambda: None)
    with TestClient(app, base_url="http://127.0.0.1") as reopened:
        state = reopened.get("/api/state").json()
        assert state["documents"][0]["id"] == document["id"]
        assert state["token"] != client.headers["X-Workspace-Token"]
        reopened.headers["X-Workspace-Token"] = state["token"]
        assert not ask(reopened, [document["id"]]).json()["abstained"]


def test_pdf_page_citation_and_original_file(client):
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 120), NOTES.split("\n\n")[1], fontname="china-s", fontsize=11)
        content = pdf.tobytes()
    response = client.post("/api/documents?filename=lesson.pdf", content=content)
    assert response.status_code == 201, response.text
    result = ask(client, [response.json()["id"]]).json()
    assert not result["abstained"]
    source = result["sources"][0]
    assert source["page"] == 1 and source["url"].endswith("#page=1")
    original = client.get(source["url"])
    assert original.content == content
    assert original.headers["content-type"] == "application/pdf"
    assert "inline" in original.headers["content-disposition"]


def test_pptx_upload_cites_slide(client):
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "训练集与测试集"
    slide.placeholders[1].text = NOTES.split("\n\n")[1]
    buffer = io.BytesIO()
    slides.save(buffer)
    response = client.post("/api/documents?filename=lesson.pptx", content=buffer.getvalue())
    assert response.status_code == 201, response.text
    result = ask(client, [response.json()["id"]]).json()
    assert not result["abstained"]
    assert result["sources"][0]["page"] == 1 and result["sources"][0]["type"] == "pptx"


@pytest.mark.parametrize("name,content,status", [
    ("notes.txt", b"text", 415), ("notes.md", b"", 413),
    ("notes.pdf", b"not a PDF", 422), ("notes.pptx", b"not a PPTX", 422),
    ("empty.md", b"# heading only", 422),
])
def test_invalid_uploads_leave_no_saved_documents(client, name, content, status):
    response = client.post("/api/documents", params={"filename": name}, content=content)
    assert response.status_code == status
    assert client.get("/api/state").json()["documents"] == []
    assert not list(client.app.state.library.root.iterdir())


def test_upload_limit_is_enforced_on_received_body(client, monkeypatch):
    monkeypatch.setattr("course_rag_assistant.web.MAX_UPLOAD", 4)
    assert client.post("/api/documents?filename=large.md", content=b"12345").status_code == 413


def test_filename_cannot_escape_library(client):
    document = upload(client, "../../escaped.md")
    assert document["name"] == "escaped.md"
    stored = client.app.state.library.documents[document["id"]]
    path = client.app.state.library.root / stored.id / stored.source
    assert path.is_file() and path.resolve().is_relative_to(client.app.state.library.root)
    assert client.get("/api/documents/not-an-id/file").status_code == 404


def test_foreign_origin_missing_token_and_foreign_host_rejected(client):
    assert client.post("/api/demo", headers={"X-Workspace-Token": "wrong"}).status_code == 403
    assert client.post("/api/demo", headers={"Origin": "https://foreign.example"}).status_code == 403
    assert client.get("/api/state", headers={"Host": "foreign.example"}).status_code == 400


def test_empty_or_unknown_scope_and_empty_question(client):
    assert ask(client, []).status_code == 422
    assert ask(client, ["unknown"]).status_code == 404
    doc = upload(client)
    assert ask(client, [doc["id"]], "   ").status_code == 400
    assert ask(client, [doc["id"]], "a" * 1001).status_code == 422


def test_unrelated_question_has_no_fake_citations(client):
    doc = upload(client)
    result = ask(client, [doc["id"]], "护照如何办理？").json()
    assert result["abstained"] and not result["sources"]
    assert all(not c["source_ids"] for c in result["claims"])


def test_demo_is_idempotent_and_static_page_available(client):
    a = client.post("/api/demo").json()
    b = client.post("/api/demo").json()
    assert a["id"] == b["id"]
    assert len(client.get("/api/state").json()["documents"]) == 1
    assert client.get("/").status_code == 200
    assert client.get("/assets/app.js").status_code == 200
    assert client.get("/assets/style.css").status_code == 200


def test_unconfigured_model_does_not_silently_claim_model_answer(client):
    doc = upload(client)
    assert not client.get("/api/state").json()["model_configured"]
    assert ask(client, [doc["id"]], use_model=True).status_code == 409


@pytest.mark.parametrize("broken", [False, True])
def test_model_selection_and_failure_fallback_are_reported(tmp_path, broken):
    class Generator:
        def generate(self, prompt):
            if broken:
                raise RuntimeError("local service unavailable")
            return '{"passage_ids":[1]}' if "passage_ids" in prompt else '{"sentence_ids":[1]}'

    app = create_app(tmp_path / "library", generator_factory=Generator)
    with TestClient(app, base_url="http://127.0.0.1") as client:
        client.headers["X-Workspace-Token"] = client.get("/api/state").json()["token"]
        doc = upload(client)
        result = ask(client, [doc["id"]], use_model=True).json()
        assert result["mode"] == ("fallback" if broken else "model")
        assert not result["abstained"] and result["sources"]
